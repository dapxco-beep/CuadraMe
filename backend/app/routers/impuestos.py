"""Impuestos (Colombia): contador de 4x1000 y organizador de posibles deducibles.
Solo informativo. La lógica vive en app/impuestos.py (funciones puras)."""
import io
from datetime import datetime, timedelta
from typing import Literal, Optional

from fastapi import APIRouter, Depends, Header, HTTPException, Query
from fastapi.responses import Response, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from app import renta_alerta
from app.auth import UsuarioActual, cron_autorizado, get_current_user
from app.crons import Corrida, liberar_marca, reclamar_marca
from app.compartidos import cuenta_accesible, cuentas_ids_accesibles
from app.contabilidad import tipos_cuenta_de, tipos_de_cuentas, tipos_efectivos
from app.database import get_db
from app.fx import convertir_a_cop
from app.impuestos import AVISO, CATEGORIAS_FISCALES, calcular_gmf, gmf_mes, organizar_deducibles
from app.models import Cuenta, Movimiento
from app.models_fiscal import AjusteUsuario, MarcaFiscal
from app.utiles import ahora_co, celda_segura

router = APIRouter(prefix="/impuestos", tags=["impuestos"])

_TRES_DIAS = timedelta(days=3)  # margen para emparejar traslados del borde del año


class AjustesFiscalesUpdate(BaseModel):
    cuenta_exenta_gmf_id: Optional[str] = None


class MarcaFiscalUpdate(BaseModel):
    categoria_fiscal: str = Field(max_length=40)
    soporte: str = Field(default="", max_length=300)
    estado: Literal["confirmada", "descartada"] = "confirmada"


def _anio_valido(anio: int | None) -> int:
    anio = anio or ahora_co().year
    if not 2000 <= anio <= ahora_co().year + 1:
        raise HTTPException(status_code=422, detail="Año fuera de rango")
    return anio


def _datos_anio(db: Session, usuario: UsuarioActual, anio: int) -> dict:
    ids = cuentas_ids_accesibles(db, usuario.id)
    cuentas = db.query(Cuenta).filter(Cuenta.id.in_(ids)).all()
    monedas = {c.id: c.moneda for c in cuentas}
    tipos = tipos_de_cuentas(cuentas)
    movs = (
        db.query(Movimiento)
        .filter(Movimiento.cuenta_id.in_(ids), Movimiento.fecha >= datetime(anio, 1, 1) - _TRES_DIAS,
                Movimiento.fecha < datetime(anio + 1, 1, 1))
        .all()
    )
    te = tipos_efectivos(movs, tipos)
    en_anio = [m for m in movs if m.fecha.year == anio]

    def cop(m):
        return convertir_a_cop(m.monto, monedas.get(m.cuenta_id, "COP"))

    gmf_filas = [
        {"fecha": m.fecha, "monto": cop(m), "tipo": m.tipo, "fuente": m.fuente, "notas": m.notas or "",
         "descripcion": m.descripcion, "cuenta_id": m.cuenta_id}
        for m in en_anio
    ]
    ajuste = db.query(AjusteUsuario).filter(AjusteUsuario.usuario_id == usuario.id).first()
    exenta = ajuste.cuenta_exenta_gmf_id if ajuste else None
    gmf = calcular_gmf(gmf_filas, tipos, exenta, anio)

    marcas = {
        r.movimiento_id: {"categoria_fiscal": r.categoria_fiscal, "soporte": r.soporte or "", "estado": r.estado}
        for r in db.query(MarcaFiscal).filter(MarcaFiscal.usuario_id == usuario.id).all()
    }
    gastos = [
        {"id": m.id, "fecha": m.fecha, "descripcion": m.descripcion, "categoria": m.categoria, "monto": cop(m)}
        for m in en_anio if te[m.id] == "expense"
    ]
    deducibles = organizar_deducibles(gastos, marcas)
    for i in deducibles["items"]:
        i["fecha"] = i["fecha"].strftime("%Y-%m-%d")
    cuentas_ahorro = [{"id": c.id, "nombre": c.nombre} for c in cuentas if c.tipo == "ahorro"]
    return {"anio": anio, "aviso": AVISO, "gmf": gmf, "deducibles": deducibles,
            "cuenta_exenta_gmf_id": exenta, "cuentas_ahorro": cuentas_ahorro,
            "categorias_fiscales": {k: {"etiqueta": v[0], "ayuda": v[1]} for k, v in CATEGORIAS_FISCALES.items()}}


@router.get("/resumen")
def resumen(
    anio: int | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario: UsuarioActual = Depends(get_current_user),
):
    return _datos_anio(db, usuario, _anio_valido(anio))


@router.put("/ajustes")
def guardar_ajustes(
    data: AjustesFiscalesUpdate,
    db: Session = Depends(get_db),
    usuario: UsuarioActual = Depends(get_current_user),
):
    """Marca cuál es tu cuenta exenta del 4x1000 (la ley permite una por titular)."""
    if data.cuenta_exenta_gmf_id:
        c = cuenta_accesible(db, usuario.id, data.cuenta_exenta_gmf_id)
        if not c or c.usuario_id != usuario.id or c.tipo != "ahorro":
            raise HTTPException(status_code=404, detail="Elige una de tus cuentas de ahorro")
    ajuste = db.query(AjusteUsuario).filter(AjusteUsuario.usuario_id == usuario.id).first()
    if not ajuste:
        ajuste = AjusteUsuario(usuario_id=usuario.id)
        db.add(ajuste)
    ajuste.cuenta_exenta_gmf_id = data.cuenta_exenta_gmf_id or None
    db.commit()
    return {"ok": True, "cuenta_exenta_gmf_id": ajuste.cuenta_exenta_gmf_id}


@router.put("/marcas/{movimiento_id}")
def marcar(
    movimiento_id: str,
    data: MarcaFiscalUpdate,
    db: Session = Depends(get_db),
    usuario: UsuarioActual = Depends(get_current_user),
):
    """Confirma (o descarta) un movimiento como posible deducible y guarda el soporte como nota."""
    if data.categoria_fiscal not in CATEGORIAS_FISCALES:
        raise HTTPException(status_code=422, detail="Categoría fiscal desconocida")
    mov = db.query(Movimiento).filter(Movimiento.id == movimiento_id).first()
    if not mov or mov.cuenta_id not in cuentas_ids_accesibles(db, usuario.id):
        raise HTTPException(status_code=404, detail="Movimiento no encontrado")
    marca = (
        db.query(MarcaFiscal)
        .filter(MarcaFiscal.usuario_id == usuario.id, MarcaFiscal.movimiento_id == movimiento_id)
        .first()
    )
    if not marca:
        marca = MarcaFiscal(usuario_id=usuario.id, movimiento_id=movimiento_id)
        db.add(marca)
    marca.categoria_fiscal = data.categoria_fiscal
    marca.soporte = data.soporte.strip()
    marca.estado = data.estado
    db.commit()
    return {"ok": True}


@router.delete("/marcas/{movimiento_id}")
def quitar_marca(
    movimiento_id: str,
    db: Session = Depends(get_db),
    usuario: UsuarioActual = Depends(get_current_user),
):
    """Vuelve el movimiento a su estado original (el sistema puede volver a sugerirlo)."""
    db.query(MarcaFiscal).filter(
        MarcaFiscal.usuario_id == usuario.id, MarcaFiscal.movimiento_id == movimiento_id
    ).delete()
    db.commit()
    return {"ok": True}


@router.get("/exportar.xlsx")
def exportar_excel(
    anio: int | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario: UsuarioActual = Depends(get_current_user),
):
    """Excel para organizar la declaración de renta: resumen, 4x1000 por mes y detalle de deducibles."""
    d = _datos_anio(db, usuario, _anio_valido(anio))
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.append([f"CuadraMe — apoyo para la declaración de renta {d['anio']}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([d["aviso"]])
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:E2")
    ws.row_dimensions[2].height = 60
    ws.append([])
    ws.append(["4x1000 estimado del año", d["gmf"]["total_estimado"]])
    ws.append(["4x1000 registrado (líneas 'gravamen')", d["gmf"]["total_registrado"]])
    ws.append(["UVT usada", d["gmf"]["uvt"]])
    ws.append([])
    ws.append(["Posibles deducibles", "Confirmado por ti", "Sugerido sin revisar", "Cantidad"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for cat in d["deducibles"]["por_categoria"]:
        ws.append([cat["etiqueta"], cat["confirmado"], cat["sugerido"], cat["cantidad"]])
    ws.column_dimensions["A"].width = 42
    for col in "BCDE":
        ws.column_dimensions[col].width = 22

    g = wb.create_sheet("4x1000 por mes")
    g.append(["Mes", "Débitos", "Base gravable", "Exento usado", "GMF estimado", "GMF registrado"])
    for c in g[1]:
        c.font = Font(bold=True)
    for f in d["gmf"]["meses"]:
        g.append([f["mes"], f["debitos"], f["base_gravable"], f["exento_usado"], f["gmf_estimado"], f["gmf_registrado"]])
    for col in "ABCDEF":
        g.column_dimensions[col].width = 18

    det = wb.create_sheet("Posibles deducibles")
    det.append(["Fecha", "Descripción", "Monto", "Categoría fiscal", "Estado", "Soporte / nota"])
    for c in det[1]:
        c.font = Font(bold=True)
    for i in d["deducibles"]["items"]:
        det.append([i["fecha"], celda_segura(i["descripcion"]), i["monto"], CATEGORIAS_FISCALES.get(i["categoria_fiscal"], ("Otro",))[0],
                    "Confirmado por ti" if i["estado"] == "confirmado" else "Sugerido (sin revisar)", celda_segura(i["soporte"])])
    det.column_dimensions["A"].width = 12
    det.column_dimensions["B"].width = 44
    det.column_dimensions["C"].width = 16
    det.column_dimensions["D"].width = 26
    det.column_dimensions["E"].width = 24
    det.column_dimensions["F"].width = 44

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=cuadrame_renta_{d['anio']}.xlsx"},
    )


# ------------------------------------------------------------ ¿Tengo que declarar renta? y 4x1000 del mes

def _anio_gravable(anio: int | None) -> int:
    hoy = ahora_co()
    anio = anio or renta_alerta.anio_gravable_vigente(hoy.date())
    if not 2000 <= anio <= hoy.year:
        raise HTTPException(status_code=422, detail="Año fuera de rango")
    return anio


@router.get("/renta-estimada")
def renta_estimada(
    anio: int | None = Query(default=None, description="Año gravable (por defecto el anterior al actual)"),
    db: Session = Depends(get_db),
    usuario: UsuarioActual = Depends(get_current_user),
):
    """Los cinco topes que obligan a declarar renta frente a tus datos. Informativo, no es asesoría."""
    return renta_alerta.renta_de_usuario(db, usuario.id, _anio_gravable(anio), ahora_co())


@router.get("/renta-resumen.csv")
def renta_resumen_csv(
    anio: int | None = Query(default=None),
    db: Session = Depends(get_db),
    usuario: UsuarioActual = Depends(get_current_user),
):
    """Totales del año por criterio (sin movimientos individuales) para llevar al contador o a la DIAN."""
    r = renta_alerta.renta_de_usuario(db, usuario.id, _anio_gravable(anio), ahora_co())
    return Response(
        content="﻿" + renta_alerta.resumen_csv(r), media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename=cuadrame_resumen_renta_{r['anio']}.csv"},
    )


@router.get("/gmf-estimado")
def gmf_estimado(
    mes: str | None = Query(default=None, pattern=r"^\d{4}-(0[1-9]|1[0-2])$", description="YYYY-MM (por defecto el mes actual)"),
    db: Session = Depends(get_db),
    usuario: UsuarioActual = Depends(get_current_user),
):
    """4x1000 pagado en el mes y qué tan cerca estás del tope exento de tu cuenta marcada."""
    mes = mes or ahora_co().strftime("%Y-%m")
    anio, m = int(mes[:4]), int(mes[5:])
    if not 2000 <= anio <= ahora_co().year + 1:
        raise HTTPException(status_code=422, detail="Mes fuera de rango")
    ids = cuentas_ids_accesibles(db, usuario.id)
    cuentas = db.query(Cuenta).filter(Cuenta.id.in_(ids)).all()
    monedas = {c.id: c.moneda for c in cuentas}
    inicio = datetime(anio, m, 1)
    fin = datetime(anio + (m == 12), m % 12 + 1, 1)
    movs = db.query(Movimiento).filter(Movimiento.cuenta_id.in_(ids), Movimiento.fecha >= inicio,
                                       Movimiento.fecha < fin).all() if ids else []
    filas = [{"fecha": x.fecha, "monto": convertir_a_cop(x.monto, monedas.get(x.cuenta_id, "COP")), "tipo": x.tipo,
              "fuente": x.fuente, "notas": x.notas or "", "descripcion": x.descripcion,
              "categoria": x.categoria or "", "cuenta_id": x.cuenta_id} for x in movs]
    ajuste = db.query(AjusteUsuario).filter(AjusteUsuario.usuario_id == usuario.id).first()
    exenta = ajuste.cuenta_exenta_gmf_id if ajuste else None
    out = gmf_mes(filas, tipos_de_cuentas(cuentas), exenta, mes)
    out["cuentas_ahorro"] = [{"id": c.id, "nombre": c.nombre} for c in cuentas if c.tipo == "ahorro" and c.usuario_id == usuario.id]
    return out


CLAVE_AVISO_RENTA = "aviso_renta"
HITOS_AVISO_RENTA = (14, 3)  # días antes del plazo


def hito_aviso(dias_restantes: int) -> int | None:
    """Cuál aviso toca hoy: 3 si faltan <= 3 días, 14 si faltan <= 14, nada si falta más o ya venció."""
    if dias_restantes < 0:
        return None
    for h in sorted(HITOS_AVISO_RENTA):
        if dias_restantes <= h:
            return h
    return None


@router.post("/cron/avisar-renta")
def cron_avisar_renta(db: Session = Depends(get_db), x_cron_secret: str = Header(default="")):
    """Diario. Push a quienes (con push activo) están en «sí» o «cerca» cuando faltan 14 y 3 días.
    Idempotente por usuario, año y hito (marcas_cron). Tolerante a fallos por usuario."""
    if not cron_autorizado(x_cron_secret):
        raise HTTPException(status_code=401, detail="No autorizado")
    from app.models import PushSubscription
    from app.push import enviar_push

    hoy = ahora_co()
    anio = renta_alerta.anio_gravable_vigente(hoy.date())
    plazo = renta_alerta.plazo_del_anio(anio, hoy.date())
    res = {"ok": True, "anio": anio, "enviados": 0, "omitidos": 0}
    if not plazo or not plazo["abierto"] or hito_aviso(plazo["dias_restantes"]) is None:
        return {**res, "motivo": "fuera_de_ventana", "procesados": 0, "fallidos": 0}
    hito = hito_aviso(plazo["dias_restantes"])
    uids = [u for (u,) in db.query(PushSubscription.usuario_id).distinct().all()]
    corrida = Corrida(db)
    for uid in uids:
        with corrida.aislar(uid):
            if not reclamar_marca(db, uid, f"{CLAVE_AVISO_RENTA}:{anio}", str(hito)):
                continue
            try:
                r = renta_alerta.renta_de_usuario(db, uid, anio, hoy)
                if r["resultado"]["clave"] not in ("si", "cerca"):
                    res["omitidos"] += 1
                    continue
                dias = plazo["dias_restantes"]
                cuando = "hoy" if dias == 0 else f"en {dias} días"
                n = enviar_push(db, uid, "¿Tienes que declarar renta?",
                                f"El plazo vence {cuando} (26 de octubre). Revisa si te toca. Es orientación, no asesoría.",
                                "/impuestos", boton="Ver mi resultado")
                res["enviados"] += 1 if n else 0
            except Exception:
                liberar_marca(db, uid, f"{CLAVE_AVISO_RENTA}:{anio}")
                raise
    return {**res, **corrida.resumen()}

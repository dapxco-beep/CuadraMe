import csv
import io
from datetime import datetime

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.auth import UsuarioActual, get_current_user
from app.compartidos import cuentas_ids_accesibles
from app.contabilidad import delta_saldo, tipos_cuenta_de, tipos_efectivos
from app.database import get_db
from app.models import Cuenta, Movimiento
from app.utiles import celda_segura

router = APIRouter(prefix="/exportar", tags=["exportar"])


def _movimientos_rango(
    db: Session, usuario_id: str, desde: str | None, hasta: str | None
) -> list[Movimiento]:
    ids_accesibles = cuentas_ids_accesibles(db, usuario_id)
    query = db.query(Movimiento).filter(Movimiento.cuenta_id.in_(ids_accesibles))
    if desde:
        query = query.filter(Movimiento.fecha >= datetime.fromisoformat(desde))
    if hasta:
        query = query.filter(Movimiento.fecha <= datetime.fromisoformat(hasta))
    return query.order_by(Movimiento.fecha.asc()).all()


ETIQUETA_CONTABLE = {"expense": "Gasto", "income": "Ingreso", "transfer": "No cuenta (pago de tarjeta o traslado propio)"}


def _etiquetas_contables(db: Session, usuario_id: str, movimientos: list[Movimiento]) -> dict[str, str]:
    """Qué papel juega cada fila en los resúmenes: así quien suma el Excel por su cuenta
    no cuenta dos veces las compras con tarjeta y el pago de esa tarjeta."""
    te = tipos_efectivos(movimientos, tipos_cuenta_de(db, cuentas_ids_accesibles(db, usuario_id)))
    return {mid: ETIQUETA_CONTABLE[t] for mid, t in te.items()}


@router.get("/movimientos.csv")
def exportar_csv(
    desde: str | None = None,
    hasta: str | None = None,
    db: Session = Depends(get_db),
    usuario: UsuarioActual = Depends(get_current_user),
):
    movimientos = _movimientos_rango(db, usuario.id, desde, hasta)

    buffer = io.StringIO()
    writer = csv.writer(buffer)
    etiquetas = _etiquetas_contables(db, usuario.id, movimientos)
    writer.writerow(["Fecha", "Descripción", "Categoría", "Tipo", "Monto", "Fuente", "Notas", "Cuenta en resúmenes como", "Texto del banco", "Subcategoría"])
    for m in movimientos:
        writer.writerow([
            m.fecha.strftime("%Y-%m-%d"), celda_segura(m.titulo or m.descripcion), celda_segura(m.categoria),
            m.tipo, m.monto, m.fuente, celda_segura(m.notas), etiquetas[m.id],
            celda_segura(m.descripcion), celda_segura(m.subcategoria or ""),
        ])

    buffer.seek(0)
    return StreamingResponse(
        iter([buffer.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=movimientos_cuadrame.csv"},
    )


TIPO_ES = {"expense": "Gasto", "income": "Ingreso", "transfer": "Traslado"}
FORMATO_PESOS = '"$" #,##0;[Red]-"$" #,##0'
FORMATO_PESOS_CENT = '"$" #,##0.00;[Red]-"$" #,##0.00'
MESES_ES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]
COLOR_ENCABEZADO = "1F3A5F"


def _encabezar(ws, fila: int = 1) -> None:
    for c in ws[fila]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=COLOR_ENCABEZADO)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[fila].height = 28


def _anchos(ws, tope: int = 48, minimo: int = 10) -> None:
    for col in ws.columns:
        largo = max(len(str(c.value if not hasattr(c.value, "strftime") else "00/00/0000")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = max(minimo, min(tope, largo + 3))


@router.get("/movimientos.xlsx")
def exportar_excel(
    desde: str | None = None,
    hasta: str | None = None,
    db: Session = Depends(get_db),
    usuario: UsuarioActual = Depends(get_current_user),
):
    """Excel listo para usar: fechas y montos con formato real (no texto), tipos en español, la cuenta de cada movimiento,
    filtros y primera fila fija, una hoja de resumen por mes y por categoría, y una hoja con el estado de las cuentas.
    Los resúmenes salen calculados (no fórmulas) para que se vean igual en el celular y en cualquier visor."""
    movimientos = _movimientos_rango(db, usuario.id, desde, hasta)
    cuentas = {c.id: c for c in db.query(Cuenta).filter(Cuenta.id.in_(cuentas_ids_accesibles(db, usuario.id))).all()}
    etiquetas = _etiquetas_contables(db, usuario.id, movimientos)
    efectivos = tipos_efectivos(movimientos, tipos_cuenta_de(db, cuentas_ids_accesibles(db, usuario.id)))

    wb = Workbook()
    ws = wb.active
    ws.title = "Movimientos"
    encabezados = ["Fecha", "Descripción", "Categoría", "Subcategoría", "Cuenta", "Tipo", "Monto", "Valor (+ entra / − sale)",
                   "Cuenta en resúmenes como", "Fuente", "Notas", "Texto del banco"]
    ws.append(encabezados)
    for m in movimientos:
        cuenta = cuentas.get(m.cuenta_id)
        ws.append([
            m.fecha.date(), celda_segura(m.titulo or m.descripcion), celda_segura(m.categoria), celda_segura(m.subcategoria or ""),
            celda_segura(cuenta.nombre if cuenta else ""), TIPO_ES.get(m.tipo, m.tipo), float(m.monto), float(delta_saldo(m)),
            etiquetas[m.id], m.fuente, celda_segura(m.notas), celda_segura(m.descripcion),
        ])
    _encabezar(ws)
    for fila in ws.iter_rows(min_row=2):
        fila[0].number_format = "dd/mm/yyyy"
        fila[0].alignment = Alignment(horizontal="center")
        fila[6].number_format = FORMATO_PESOS_CENT
        fila[7].number_format = FORMATO_PESOS_CENT
    ws.freeze_panes = "A2"
    if movimientos:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(encabezados))}{len(movimientos) + 1}"
    _anchos(ws)

    # Hoja «Resumen»: ingresos y gastos por mes y gastos por categoría, con la MISMA regla que la app
    # (pagos de tarjeta y traslados propios no cuentan: no se suman dos veces).
    por_mes: dict[tuple[int, int], list[float]] = {}
    por_categoria: dict[str, float] = {}
    for m in movimientos:
        t = efectivos.get(m.id)
        if t not in ("income", "expense"):
            continue
        fila = por_mes.setdefault((m.fecha.year, m.fecha.month), [0.0, 0.0])
        fila[0 if t == "income" else 1] += float(m.monto)
        if t == "expense":
            por_categoria[m.categoria or "Sin categoría"] = por_categoria.get(m.categoria or "Sin categoría", 0.0) + float(m.monto)
    rs = wb.create_sheet("Resumen")
    rs.append(["Mes", "Ingresos", "Gastos", "Neto (ingresos − gastos)"])
    _encabezar(rs)
    for (anio, mes), (ing, gas) in sorted(por_mes.items()):
        rs.append([f"{MESES_ES[mes - 1].capitalize()} {anio}", round(ing, 2), round(gas, 2), round(ing - gas, 2)])
    if por_mes:
        rs.append(["Total", round(sum(v[0] for v in por_mes.values()), 2), round(sum(v[1] for v in por_mes.values()), 2),
                   round(sum(v[0] - v[1] for v in por_mes.values()), 2)])
        for c in rs[rs.max_row]:
            c.font = Font(bold=True)
    rs.append([])
    inicio_cat = rs.max_row + 1
    rs.append(["Gastos por categoría", "Total", "% del gasto"])
    _encabezar(rs, inicio_cat)
    total_gasto = sum(por_categoria.values()) or 1.0
    for cat, total in sorted(por_categoria.items(), key=lambda kv: -kv[1]):
        rs.append([celda_segura(cat), round(total, 2), round(total / total_gasto, 4)])
        rs.cell(rs.max_row, 3).number_format = "0.0%"
    for fila in rs.iter_rows(min_row=2):
        for c in fila[1:]:
            if c.number_format == "General":
                c.number_format = FORMATO_PESOS
    rs.freeze_panes = "A2"
    _anchos(rs)

    # Hoja «Cuentas»: lo que hay hoy en cada cuenta (una tarjeta es deuda, no saldo).
    cs = wb.create_sheet("Cuentas")
    cs.append(["Cuenta", "Tipo", "Saldo / deuda hoy", "Cupo"])
    _encabezar(cs)
    for c in sorted(cuentas.values(), key=lambda c: (c.tipo == "credito", c.nombre or "")):
        es_tarjeta = c.tipo == "credito"
        cs.append([celda_segura(c.nombre), "Tarjeta de crédito (deuda)" if es_tarjeta else (c.tipo or "").capitalize(),
                   round(-float(c.saldo or 0), 2) if es_tarjeta else round(float(c.saldo or 0), 2),
                   float(c.cupo_total or 0) or None])
    for fila in cs.iter_rows(min_row=2):
        fila[2].number_format = FORMATO_PESOS
        fila[3].number_format = FORMATO_PESOS
    cs.freeze_panes = "A2"
    _anchos(cs)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=movimientos_cuadrame.xlsx"},
    )

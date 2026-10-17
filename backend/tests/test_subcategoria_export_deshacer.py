"""Subcategoría: el CSV/Excel la trae (columna al final, sin mover las anteriores) y «Deshacer» una recategorización
que solo agregó subcategoría (misma categoría) la quita."""
import csv
import io
import os
import tempfile

os.environ.setdefault("DATABASE_URL", f"sqlite:///{os.path.join(tempfile.mkdtemp(), 'sub_exp.db')}")

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.auth import UsuarioActual, get_current_user  # noqa: E402
from app.main import app  # noqa: E402

client = TestClient(app)


@pytest.fixture(autouse=True)
def _auth():
    app.dependency_overrides[get_current_user] = lambda: UsuarioActual(id="sub-exp", email="s@test.com", nombre="S")
    yield
    app.dependency_overrides.pop(get_current_user, None)


def _preparar():
    cuenta = client.post("/cuentas", json={"nombre": "Banco", "saldo": 0}).json()
    client.post("/subcategorias", json={"categoria": "Comida", "nombre": "Restaurantes"})  # 400 si ya existe (otra prueba)
    mov = client.post("/movimientos", json={"cuenta_id": cuenta["id"], "descripcion": "ALMUERZO", "monto": 20000, "tipo": "expense",
                                            "categoria": "Comida"}).json()
    return cuenta, mov


def _sub(mov_id):
    return next(m for m in client.get("/movimientos").json() if m["id"] == mov_id)["subcategoria"]


def test_csv_y_excel_traen_la_subcategoria_al_final():
    _, mov = _preparar()
    client.post("/movimientos/recategorizar-masivo", json={"movimiento_ids": [mov["id"]], "categoria": "Comida", "subcategoria": "Restaurantes"})
    filas = list(csv.reader(io.StringIO(client.get("/exportar/movimientos.csv").text)))
    enc = filas[0]
    assert enc[:9] == ["Fecha", "Descripción", "Categoría", "Tipo", "Monto", "Fuente", "Notas", "Cuenta en resúmenes como", "Texto del banco"]
    assert enc[-1] == "Subcategoría"
    fila = next(f for f in filas if "ALMUERZO" in f)
    assert fila[-1] == "Restaurantes" and fila[8] == "ALMUERZO"
    ws = load_workbook(io.BytesIO(client.get("/exportar/movimientos.xlsx").content)).active
    assert [c.value for c in ws[1]][-1] == "Subcategoría"
    assert "Restaurantes" in [c.value for c in ws[2]]


def test_deshacer_una_recategorizacion_que_solo_agrego_subcategoria_la_quita():
    _, mov = _preparar()
    assert _sub(mov["id"]) in (None, "")
    client.post("/movimientos/recategorizar-masivo", json={"movimiento_ids": [mov["id"]], "categoria": "Comida", "subcategoria": "Restaurantes"})
    assert _sub(mov["id"]) == "Restaurantes"
    # Sin el campo, la misma categoría conserva la subcategoría (comportamiento de siempre).
    client.post("/movimientos/recategorizar-masivo", json={"movimiento_ids": [mov["id"]], "categoria": "Comida"})
    assert _sub(mov["id"]) == "Restaurantes"
    # Deshacer: la misma categoría con subcategoria "" la limpia.
    client.post("/movimientos/recategorizar-masivo", json={"movimiento_ids": [mov["id"]], "categoria": "Comida", "subcategoria": ""})
    assert _sub(mov["id"]) in (None, "")

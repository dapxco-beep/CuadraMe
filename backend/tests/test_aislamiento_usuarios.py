"""
Auditoría de seguridad: AISLAMIENTO entre usuarios (IDOR) en TODOS los routers.

Se crean dos usuarios sintéticos. A es el dueño y se siembra con datos en cada
módulo (cuentas, movimientos, correos, extractos, reglas, metas...). B es el
atacante: intenta leer, modificar y borrar los datos de A usando sus ids.

Qué se verifica:
  1. Ningún endpoint de LISTA/lectura de B devuelve datos de A (marcador ZZA).
  2. Cada endpoint con id de recurso (o con ids de A dentro del cuerpo) responde
     4xx para B y NO modifica nada de A (se compara una foto completa de las filas de A).
  3. Todo endpoint del backend está clasificado: si alguien agrega uno nuevo y no lo
     cubre aquí, la prueba `test_todo_endpoint_esta_clasificado` falla.
  4. Sin sesión, todo endpoint de usuario responde 401; los crons exigen su secreto.
"""
import importlib
import json
import os
import pkgutil
import tempfile
from datetime import datetime

os.environ.setdefault("DATABASE_URL", f"sqlite:///{os.path.join(tempfile.mkdtemp(), 'aislamiento.db')}")

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

import app.routers as paquete_routers  # noqa: E402
from app.auth import UsuarioActual, get_current_user  # noqa: E402
from app.database import Base, SessionLocal  # noqa: E402
from app.main import app  # noqa: E402
from app.models import CorreoGasto, ExtractoCorreo, HistorialExtracto, PushSubscription  # noqa: E402

client = TestClient(app)
A = UsuarioActual(id="aisl-a", email="a-aisl@test.com", nombre="Dueña A")
B = UsuarioActual(id="aisl-b", email="b-aisl@test.com", nombre="Atacante B")
C = UsuarioActual(id="aisl-c", email="c-aisl@test.com", nombre="Invitado C")
MARCA = "ZZA"  # aparece en todo texto sembrado para A
ACTUAL = {"u": A}
I: dict = {}  # ids sembrados de A


@pytest.fixture(autouse=True)
def _auth():
    app.dependency_overrides[get_current_user] = lambda: ACTUAL["u"]
    yield
    app.dependency_overrides.pop(get_current_user, None)


@pytest.fixture(autouse=True, scope="module")
def _limpiar_al_final():
    """Las pruebas comparten base con el resto de la suite: no dejar filas sembradas."""
    yield
    from app.derechos_titular import borrar_todo

    db = SessionLocal()
    try:
        for u in (A, B, C):
            borrar_todo(db, u.id, u.email)
    finally:
        db.close()


def como(u):
    ACTUAL["u"] = u


def _ok(r):
    assert r.status_code in (200, 201), f"{r.request.method} {r.request.url.path} -> {r.status_code}: {r.text[:300]}"
    return r.json()


def _sembrar():
    """Siembra datos de A en todos los módulos (idempotente: corre una sola vez)."""
    if I:
        return
    from cryptography.fernet import Fernet

    from app.config import settings

    if not settings.fernet_key:
        settings.fernet_key = Fernet.generate_key().decode()
    como(A)
    banco = _ok(client.post("/cuentas", json={"nombre": f"{MARCA} Banco", "saldo": 1000}))
    otra = _ok(client.post("/cuentas", json={"nombre": f"{MARCA} Efectivo", "saldo": 500}))
    tarjeta = _ok(client.post("/cuentas", json={"nombre": f"{MARCA} Visa", "tipo": "credito", "saldo": 0,
                                                "cupo_total": 5_000_000, "dia_corte": 10, "dia_pago": 25}))
    I.update(cuenta=banco["id"], cuenta2=otra["id"], tarjeta=tarjeta["id"])
    mov = _ok(client.post("/movimientos", json={"cuenta_id": banco["id"], "descripcion": f"{MARCA} compra",
                                                "monto": 10, "tipo": "expense"}))
    mov2 = _ok(client.post("/movimientos", json={"cuenta_id": banco["id"], "descripcion": f"{MARCA} otra",
                                                 "monto": 20, "tipo": "expense"}))
    I.update(mov=mov["id"], mov2=mov2["id"])
    # decisión de traslado recordada (regla por huella) para atacarla desde B
    _ok(client.post("/traslados/decidir", json={"ids": [mov2["id"]], "decision": "real", "recordar": True}))
    I["regla_traslado"] = client.get("/traslados/reglas").json()[0]["id"]
    _ok(client.post("/movimientos/transferencia", json={"cuenta_origen_id": banco["id"],
                                                        "cuenta_destino_id": otra["id"], "monto": 5}))
    cat = _ok(client.post("/categorias", json={"nombre": f"{MARCA}Cat", "icono": "🐶"}))
    I["categoria"] = cat["id"]
    I["regla"] = _ok(client.post("/reglas-categoria", json={"patron": "zzapatron", "categoria": "Comida"}))["id"]
    I["presupuesto"] = _ok(client.post("/presupuestos", json={"categoria": "Comida", "limite": 100}))["id"]
    I["presupuesto2"] = _ok(client.post("/presupuestos", json={"categoria": "Transporte", "limite": 100}))["id"]
    I["meta"] = _ok(client.post("/metas", json={"nombre": f"{MARCA} Meta", "monto_objetivo": 1000}))["id"]
    I["recurrente"] = _ok(client.post("/recurrentes", json={"nombre": f"{MARCA} Arriendo", "monto": 50,
                                                            "cuenta_id": banco["id"]}))["id"]
    I["deuda"] = _ok(client.post("/deudas", json={"nombre": f"{MARCA} Deuda", "saldo_inicial": 1000,
                                                  "cuota_mensual": 100}))["id"]
    I["contacto"] = _ok(client.post("/gastos-compartidos/contactos", json={"nombre": f"{MARCA} Amigo"}))["id"]
    I["grupo"] = _ok(client.post("/gastos-compartidos/grupos", json={"nombre": f"{MARCA} Viaje",
                                                                    "miembros": [I["contacto"]]}))["id"]
    gasto = _ok(client.post("/gastos-compartidos", json={
        "descripcion": f"{MARCA} cena", "total": 100, "modo": "iguales", "registro": "ninguno",
        "grupo_id": I["grupo"], "pagador_id": "yo",
        "participantes": [{"id": "yo", "valor": 1}, {"id": I["contacto"], "valor": 1}]}))
    I["gasto"] = gasto.get("id") or client.get("/gastos-compartidos").json()[0]["id"]
    liq = _ok(client.post("/gastos-compartidos/liquidaciones", json={"de_id": I["contacto"], "a_id": "yo",
                                                                    "monto": 10, "nota": MARCA}))
    I["liquidacion"] = liq.get("id") or client.get("/gastos-compartidos/liquidaciones").json()[0]["id"]
    activo = _ok(client.post("/patrimonio/activos", json={"clase": "accion", "nombre": f"{MARCA} Acción",
                                                          "unidades": 2, "costo_promedio": 10, "precio_actual": 12}))
    I["activo"] = activo["id"]
    _ok(client.post(f"/patrimonio/activos/{I['activo']}/valor", json={"valor": 13}))
    _ok(client.post(f"/patrimonio/activos/{I['activo']}/flujos", json={"monto": -20, "nota": MARCA}))
    I["flujo"] = client.get(f"/patrimonio/activos/{I['activo']}/flujos").json()[0]["id"]
    I["plan_cuotas"] = _ok(client.post("/cuotas-tarjeta", json={
        "cuenta_id": tarjeta["id"], "descripcion": f"{MARCA} nevera", "total": 1200, "num_cuotas": 12}))["id"]
    inv = _ok(client.post(f"/cuentas/{banco['id']}/invitar", json={"email": C.email, "rol": "viewer"}))
    I["invitacion"] = inv["id"]
    partida = client.post("/ahorcado/nueva-partida")
    I["partida"] = partida.json().get("id", "") if partida.status_code == 200 else ""
    I["comentario"] = _ok(client.post("/comentarios", json={"tipo": "sugerencia", "mensaje": f"{MARCA} idea",
                                                            "pantalla": "/movimientos"}))["id"]
    I["cajita"] = _ok(client.post("/rendimientos/cajitas", json={"cuenta_id": banco["id"], "nombre": f"{MARCA} Cajita",
                                                                   "saldo": 100, "tasa_ea": 9}))["id"]
    client.put("/planificacion/config", json={})
    client.put("/correos/clave-pdf", json={"clave": f"{MARCA}1234"})  # puede requerir Fernet: no es obligatorio

    db = SessionLocal()
    try:
        c = CorreoGasto(usuario_id=A.id, gmail_msg_id="aisl-msg-1", remitente=f"{MARCA}@banco.com",
                        asunto=f"{MARCA} compra", snippet=f"{MARCA} snippet", monto=99, comercio=f"{MARCA} Tienda")
        ex = ExtractoCorreo(usuario_id=A.id, gmail_msg_id="aisl-msg-2", archivo=f"{MARCA}.pdf",
                            remitente=f"{MARCA}@banco.com", asunto=f"{MARCA} extracto", banco="bancolombia")
        h = HistorialExtracto(usuario_id=A.id, cuenta_id=banco["id"], banco_detectado=f"{MARCA}banco",
                              review_json=json.dumps([{"descripcion": f"{MARCA} linea", "estado": "faltante",
                                                       "fecha": "2026-01-01", "monto": 1, "tipo": "expense",
                                                       "fingerprint": "zz"}]))
        p = PushSubscription(usuario_id=A.id, endpoint=f"https://push.example/{MARCA}", p256dh="k", auth="a")
        db.add_all([c, ex, h, p])
        db.commit()
        I.update(correo=c.id, extracto=ex.id, historial=h.id)
    finally:
        db.close()


def _foto_de(uid: str) -> dict:
    """Foto de TODAS las filas de A (tablas con usuario_id + sus cuentas y movimientos)."""
    db = SessionLocal()
    try:
        foto = {}
        for tabla in Base.metadata.sorted_tables:
            if "usuario_id" in tabla.c:
                filas = db.execute(tabla.select().where(tabla.c.usuario_id == uid)).fetchall()
                foto[tabla.name] = sorted(json.dumps(dict(f._mapping), default=str, sort_keys=True) for f in filas)
        return foto
    finally:
        db.close()


# ── 1. Lecturas: B nunca ve datos de A ──────────────────────────────────────────

def _rutas_de_routers():
    """(metodo, ruta, tiene_usuario, nombre_modulo) de cada ruta de cada router."""
    salida = []
    for m in pkgutil.iter_modules(paquete_routers.__path__):
        mod = importlib.import_module(f"app.routers.{m.name}")
        r = getattr(mod, "router", None)
        if r is None:
            continue
        for ruta in r.routes:
            usa_usuario = _depende_de(ruta.dependant, "get_current_user")
            usa_clave = _depende_de(ruta.dependant, "get_current_user_by_api_key")
            for metodo in sorted(ruta.methods - {"HEAD", "OPTIONS"}):
                salida.append((metodo, ruta.path, usa_usuario, usa_clave, m.name))
    return salida


def _depende_de(dep, nombre: str) -> bool:
    for s in dep.dependencies:
        if getattr(s.call, "__name__", "") == nombre or _depende_de(s, nombre):
            return True
    return False


RUTAS = _rutas_de_routers()
LISTAS = sorted({r for (m, r, u, k, _) in RUTAS if m == "GET" and u and "{" not in r})

# Lecturas que por diseño no devuelven datos del usuario o requieren parámetros.
LISTAS_EXENTAS = {"/correos/gmail/conectar", "/perfil/api-key", "/perfil/mis-datos"}  # mis-datos: ver test_derechos_titular.py


@pytest.mark.parametrize("ruta", [r for r in LISTAS if r not in LISTAS_EXENTAS])
def test_lecturas_de_b_no_muestran_datos_de_a(ruta):
    _sembrar()
    como(B)
    r = client.get(ruta)
    assert r.status_code < 500, f"{ruta} -> {r.status_code}"
    cuerpo = r.content.decode("latin-1") if ruta.endswith(".xlsx") else r.text
    if ruta.endswith(".xlsx"):
        import io

        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(r.content))
        cuerpo = " ".join(str(c.value) for ws in wb for fila in ws.iter_rows() for c in fila)
    assert MARCA.lower() not in cuerpo.lower(), f"{ruta} filtra datos de otro usuario"


def test_lecturas_de_a_si_ven_sus_datos():
    """Control: la marca realmente aparece para su dueña (si no, la prueba anterior no probaría nada)."""
    _sembrar()
    como(A)
    for ruta in ("/cuentas", "/movimientos", "/metas", "/gastos-compartidos/contactos", "/patrimonio/activos",
                 "/correos/lista", "/correos/extractos", "/conciliacion/historial", "/reglas-categoria",
                 "/sincronizacion/extractos"):  # «Todos los extractos»: la lista completa de A
        assert MARCA.lower() in client.get(ruta).text.lower(), ruta


# ── 2. Ataques por id: 4xx y nada cambia ─────────────────────────────────────────
# (metodo, ruta con {llaves} de I, cuerpo json, query params)
ATAQUES = [
    # traslados: B no decide sobre movimientos de A ni borra sus reglas
    ("POST", "/traslados/decidir", {"ids": ["{mov}"], "decision": "traslado", "recordar": True}),
    ("DELETE", "/traslados/reglas/{regla_traslado}", None),
    # cuentas
    ("PUT", "/cuentas/{cuenta}", {"nombre": "hackeada"}),
    ("DELETE", "/cuentas/{cuenta}", None),
    ("GET", "/cuentas/{cuenta}/detalle", None),
    ("POST", "/cuentas/{cuenta}/ajustar-saldo", {"saldo_real": 1}),
    # cuadre con el extracto: B no evalúa, ajusta, registra saldo inicial ni lee el cuadre de cuentas/historiales de A
    ("GET", "/conciliacion/cuadre?cuenta_id={cuenta}", None),
    ("GET", "/conciliacion/cuadre/que-falta?cuenta_id={cuenta}&historial_id={historial}", None),
    ("GET", "/conciliacion/cuadre/encadenado?cuenta_id={cuenta}", None),
    ("POST", "/conciliacion/cuadre/evaluar", {"cuenta_id": "{cuenta}", "historial_id": "{historial}", "saldo_final": 1, "fecha_corte": "2026-01-31"}),
    ("POST", "/conciliacion/cuadre/saldos", {"historial_id": "{historial}", "saldo_final": 1, "fecha_corte": "2026-01-31", "confirmar": True}),
    ("POST", "/conciliacion/cuadre/ajuste", {"cuenta_id": "{cuenta}", "saldo_final": 1, "fecha_corte": "2026-01-31", "confirmar": True, "confirmar_reforzado": True}),
    ("POST", "/conciliacion/cuadre/saldo-inicial", {"cuenta_id": "{cuenta}", "modo": "sin_explicar", "confirmar": True}),
    ("POST", "/conciliacion/cuadre/deshacer", {"movimiento_id": "{mov}", "confirmar": True}),
    ("POST", "/cuentas/{cuenta}/invitar", {"email": "x@test.com", "rol": "editor"}),
    ("GET", "/cuentas/{cuenta}/invitaciones", None),
    ("GET", "/cuentas/{cuenta}/compartida-con", None),
    ("DELETE", "/cuentas/{cuenta}/compartir/{usuario_c}", None),
    ("POST", "/cuentas/invitaciones/{invitacion}/aceptar", None),
    ("POST", "/cuentas/invitaciones/{invitacion}/rechazar", None),
    # movimientos (por id y con cuentas/ids de A dentro del cuerpo)
    ("GET", "/movimientos/{mov}", None),
    ("PUT", "/movimientos/{mov}", {"monto": 1}),
    ("PUT", "/movimientos/{mov}", {"cuenta_id": "{cuenta}"}),
    ("DELETE", "/movimientos/{mov}", None),
    ("POST", "/movimientos", {"cuenta_id": "{cuenta}", "descripcion": "x", "monto": 1, "tipo": "expense"}),
    ("POST", "/movimientos/transferencia", {"cuenta_origen_id": "{cuenta}", "cuenta_destino_id": "{cuenta2}", "monto": 1}),
    ("POST", "/movimientos/{mov}/recategorizar-comercio", {"categoria": "Comida"}),
    ("POST", "/movimientos/recategorizar-masivo", {"movimiento_ids": ["{mov}", "{mov2}"], "categoria": "Comida"}),
    ("POST", "/movimientos/mover-masivo", {"movimiento_ids": ["{mov}"], "cuenta_id": "{cuenta2}"}),
    ("POST", "/movimientos/eliminar-masivo", {"movimiento_ids": ["{mov}", "{mov2}"]}),
    ("POST", "/movimientos/confirmar-categorias", {"movimiento_ids": ["{mov}"]}),
    # categorías, reglas, presupuestos
    ("PUT", "/categorias/{categoria}", {"nombre": "hackeada"}),
    ("DELETE", "/categorias/{categoria}", None),
    ("DELETE", "/reglas-categoria/{regla}", None),
    ("POST", "/reglas-categoria/lote", {"ids": ["{regla}"], "accion": "borrar"}),
    ("POST", "/por-categorizar/accion", {"accion": "dejar", "movimiento_ids": ["{movimiento}"]}),
    ("POST", "/por-categorizar/deshacer", {"previos": [{"id": "{movimiento}", "categoria": "Comida"}]}),
    ("DELETE", "/presupuestos/{presupuesto}", None),
    ("POST", "/presupuestos/mover-sobre", {"presupuesto_origen_id": "{presupuesto}",
                                           "presupuesto_destino_id": "{presupuesto2}", "monto": 1}),
    # metas, recurrentes, deudas
    ("PATCH", "/metas/{meta}", {"nombre": "hackeada"}),
    ("POST", "/metas/{meta}/aportar?monto=10", None),
    ("DELETE", "/metas/{meta}", None),
    ("POST", "/metas", {"nombre": "m", "monto_objetivo": 10, "cuenta_origen_id": "{cuenta}", "aporte_automatico": 5}),
    ("POST", "/metas/{meta_propia}/aportar?monto=10&cuenta_id={cuenta}", None),
    ("POST", "/recurrentes/{recurrente}/marcar_pagado", {"crear_movimiento": True}),
    ("DELETE", "/recurrentes/{recurrente}", None),
    ("POST", "/recurrentes", {"nombre": "r", "monto": 5, "cuenta_id": "{cuenta}"}),
    ("POST", "/deudas/{deuda}/pagar", {"monto_pagado": 10, "crear_movimiento": True}),
    ("DELETE", "/deudas/{deuda}", None),
    ("POST", "/deudas", {"nombre": "d", "saldo_inicial": 10, "cuenta_pago_id": "{cuenta}"}),
    # gastos compartidos
    ("DELETE", "/gastos-compartidos/contactos/{contacto}", None),
    ("DELETE", "/gastos-compartidos/grupos/{grupo}", None),
    ("DELETE", "/gastos-compartidos/{gasto}", None),
    ("DELETE", "/gastos-compartidos/liquidaciones/{liquidacion}", None),
    ("POST", "/gastos-compartidos", {"descripcion": "x", "total": 10, "modo": "iguales", "registro": "ninguno",
                                     "grupo_id": "{grupo}", "pagador_id": "yo",
                                     "participantes": [{"id": "yo", "valor": 1}, {"id": "{contacto}", "valor": 1}]}),
    ("POST", "/gastos-compartidos", {"descripcion": "x", "total": 10, "modo": "iguales", "registro": "total",
                                     "cuenta_id": "{cuenta}", "pagador_id": "yo",
                                     "participantes": [{"id": "yo", "valor": 1}]}),
    ("POST", "/gastos-compartidos/liquidaciones", {"de_id": "{contacto}", "a_id": "yo", "monto": 5}),
    ("POST", "/gastos-compartidos/grupos", {"nombre": "g", "miembros": ["{contacto}"]}),
    ("POST", "/gastos-compartidos/calcular", {"total": 10, "modo": "iguales",
                                              "participantes": [{"id": "yo", "valor": 1}, {"id": "{contacto}", "valor": 1}]}),
    # patrimonio
    ("PUT", "/patrimonio/activos/{activo}", {"clase": "accion", "nombre": "hackeada", "unidades": 1,
                                             "costo_promedio": 1, "precio_actual": 1}),
    ("DELETE", "/patrimonio/activos/{activo}", None),
    ("POST", "/patrimonio/activos/{activo}/valor", {"valor": 1}),
    ("GET", "/patrimonio/activos/{activo}/valores", None),
    ("POST", "/patrimonio/activos/{activo}/flujos", {"monto": -1}),
    ("GET", "/patrimonio/activos/{activo}/flujos", None),
    ("DELETE", "/patrimonio/flujos/{flujo}", None),
    # tarjetas / cuotas
    ("POST", "/tarjetas/{tarjeta}/aplicar-extracto", {"cupo_total": 1}),
    ("PUT", "/cuotas-tarjeta/{plan_cuotas}", {"descripcion": "hackeada"}),
    ("POST", "/cuotas-tarjeta/{plan_cuotas}/pagar-cuota", {}),
    ("DELETE", "/cuotas-tarjeta/{plan_cuotas}", None),
    ("POST", "/cuotas-tarjeta", {"cuenta_id": "{tarjeta}", "descripcion": "x", "total": 100, "num_cuotas": 2}),
    ("GET", "/cuotas-tarjeta?cuenta_id={tarjeta}", None),
    # impuestos
    ("PUT", "/impuestos/marcas/{mov}", {"categoria_fiscal": "salud", "soporte": "x"}),
    # correos y extractos
    ("PATCH", "/correos/{correo}", {"monto": 1}),
    ("POST", "/correos/{correo}/aprobar", {"cuenta_id": "{cuenta}", "monto": 5}),
    ("POST", "/correos/{correo}/descartar", None),
    ("POST", "/correos/{correo}/no-es-duplicado", None),  # «Mostrar de todos modos» de un correo ajeno
    # buscar y unir duplicados: B no puede unir, ni marcar «no es duplicado», movimientos de A
    ("POST", "/duplicados/unir", {"conservar_id": "{mov}", "unir_ids": ["{mov2}"]}),
    ("POST", "/duplicados/no-es-duplicado", {"movimiento_ids": ["{mov}", "{mov2}"]}),
    ("POST", "/correos/extractos/{extracto}/conciliar", {"cuenta_id": "{cuenta}"}),
    ("POST", "/correos/extractos/{extracto}/descartar", None),
    # «Todos los extractos»: B no reintenta ni fija el mes de un extracto de A
    ("POST", "/correos/extractos/{extracto}/reintentar", None),
    ("PUT", "/sincronizacion/extractos/{extracto}/mes", {"mes": "2026-07"}),
    # probar la clave del PDF ahora: B no puede probar ni leer el resultado de un extracto de A
    ("POST", "/correos/clave-pdf/probar", {"extracto_id": "{extracto}"}),
    # punto de partida: B no puede fijar saldos de cuentas de A ni ajustarlas
    ("PUT", "/perfil/punto-de-partida", {"modo": "hoy", "saldos": {"{cuenta}": {"actual": 1}}, "ajustar_saldos": True}),
    ("POST", "/perfil/punto-de-partida/ajustar-saldos", {"cuenta_ids": ["{cuenta}"]}),
    # conciliación (historial de extractos)
    ("GET", "/conciliacion/historial/{historial}", None),
    ("POST", "/conciliacion/historial/{historial}/deshacer-auto", None),
    ("POST", "/conciliacion/historial/{historial}/resolver-dudosos", {"mismos": ["zz"]}),
    ("POST", "/conciliacion/historial/{historial}/importar", {"fingerprints": ["zz"]}),
    ("DELETE", "/conciliacion/historial/{historial}", None),
    ("POST", "/conciliacion/procesar", {"pdf_base64": "AAAA", "cuenta_id": "{cuenta}"}),
    # asistente / importar con cuentas de A
    ("POST", "/asistente/confirmar/movimiento", {"cuenta_id": "{cuenta}", "descripcion": "x", "monto": 5, "tipo": "expense"}),
    ("POST", "/asistente/confirmar/plan", {"nombre": "p", "monto_objetivo": 100, "aporte_mensual": 5,
                                           "cuenta_origen_id": "{cuenta}"}),
    ("POST", "/ahorcado/{partida}/adivinar", {"letra": "a"}),
    ("POST", "/importar/vista-previa", {"cuenta_id": "{cuenta}", "nombre_archivo": "a.csv", "archivo_base64": "YQ=="}),
    ("POST", "/importar/confirmar", {"cuenta_id": "{cuenta}", "filas": [{"fecha": "2026-01-01", "descripcion": "x", "monto": 5, "tipo": "expense"}]}),
    # cierre de mes: cuentas ajenas no entran al cierre de B
    ("POST", "/cierre-mes/{mes}/generar", None),
    ("DELETE", "/impuestos/marcas/{mov}", None),
    # buzón de opinión: las rutas de administración responden 404 a quien no es ADMIN_EMAIL
    ("PATCH", "/comentarios/{comentario}", {"estado": "descartado", "nota_admin": "x", "respuesta_publica": "x"}),
    ("POST", "/comentarios/marcar-vistos", {"ids": ["{comentario}"]}),
    # archivo de muestra «Mi banco no aparece»: solo el autor sube/borra; ver es solo de la administración
    ("DELETE", "/comentarios/{comentario}/adjunto", None),
    ("GET", "/comentarios/{comentario}/adjunto", None),
    ("POST", "/comentarios/{comentario}/adjunto", {}),
    # cajitas: B no puede crear cajitas en cuentas de A, ni editar, mover ni borrar las de A
    ("POST", "/rendimientos/cajitas", {"cuenta_id": "{cuenta}", "nombre": "x", "saldo": 1}),
    ("PUT", "/rendimientos/cajitas/{cajita}", {"nombre": "hackeada"}),
    ("POST", "/rendimientos/cajitas/{cajita}/mover", {"monto": 1, "sentido": "entrar"}),
    ("DELETE", "/rendimientos/cajitas/{cajita}", None),
    # tasas de referencia: edición solo de la administración (404 al resto)
    ("PUT", "/rendimientos/tasas/{cajita}", {"entidad": "x", "producto": "x", "tasa_ea": 1}),
    ("POST", "/rendimientos/tasas", {"entidad": "x", "producto": "x", "tasa_ea": 1}),
]
# Rutas donde un 200 es correcto aunque sean de B: la operación se limita a datos de B (idempotente).
# Se valida igual que la foto de A no cambie y que la respuesta no traiga la marca.
PERMITIDO_200 = {
    ("POST", "/cierre-mes/{mes}/generar"),
    ("DELETE", "/impuestos/marcas/{mov}"),  # filtra por usuario: no borra nada de A
    ("POST", "/movimientos/recategorizar-masivo"),  # ids ajenos se ignoran (0 actualizados)
    ("POST", "/movimientos/eliminar-masivo"),
    ("POST", "/movimientos/confirmar-categorias"),
    ("GET", "/cuotas-tarjeta?cuenta_id={tarjeta}"),  # lista vacía para B
    ("POST", "/gastos-compartidos/calcular"),  # solo calcula con los números que envía; no lee datos
    ("POST", "/perfil/punto-de-partida/ajustar-saldos"),  # solo mira los saldos declarados por B: ids ajenos se ignoran
    ("POST", "/conciliacion/procesar"),  # PDF inválido: debe fallar, pero no por 200 con datos de A
}


def _rellenar(valor, ids):
    if isinstance(valor, str):
        for k, v in ids.items():
            valor = valor.replace("{" + k + "}", str(v))
        return valor
    if isinstance(valor, list):
        return [_rellenar(x, ids) for x in valor]
    if isinstance(valor, dict):
        return {k: _rellenar(v, ids) for k, v in valor.items()}
    return valor


def _ids_atacante():
    como(B)
    meta_b = _ok(client.post("/metas", json={"nombre": "meta de B", "monto_objetivo": 10}))["id"]
    return {**I, "usuario_c": C.id, "meta_propia": meta_b, "mes": "2026-01"}


@pytest.mark.parametrize("metodo,ruta,cuerpo", ATAQUES, ids=[f"{m} {r}" for m, r, _ in ATAQUES])
def test_b_no_puede_tocar_recursos_de_a(metodo, ruta, cuerpo):
    _sembrar()
    ids = _ids_atacante()
    antes = _foto_de(A.id)
    como(B)
    url = _rellenar(ruta, ids)
    kw = {"json": _rellenar(cuerpo, ids)} if cuerpo is not None else {}
    r = client.request(metodo, url, **kw)
    assert r.status_code < 500, f"{metodo} {url} -> {r.status_code}: {r.text[:200]}"
    permitido = (metodo, ruta.split("?")[0]) in {(m, p.split("?")[0]) for m, p in PERMITIDO_200}
    if not permitido:
        assert 400 <= r.status_code < 500, f"B logró {metodo} {url} -> {r.status_code}: {r.text[:200]}"
    assert MARCA.lower() not in r.text.lower(), f"{metodo} {url} filtró datos de A: {r.text[:200]}"
    assert _foto_de(A.id) == antes, f"{metodo} {url} modificó datos de A"


def test_b_no_acumula_filas_con_ids_de_a():
    """Tras todos los ataques B no tiene movimientos, ni cuotas, ni nada apuntando a cuentas de A."""
    _sembrar()
    db = SessionLocal()
    try:
        from app.models import Movimiento
        from app.models_cuotas import CuotaTarjeta

        cuentas_a = {I["cuenta"], I["cuenta2"], I["tarjeta"]}
        assert db.query(Movimiento).filter(Movimiento.usuario_id == B.id,
                                           Movimiento.cuenta_id.in_(cuentas_a)).count() == 0
        assert db.query(CuotaTarjeta).filter(CuotaTarjeta.usuario_id == B.id,
                                             CuotaTarjeta.cuenta_id.in_(cuentas_a)).count() == 0
    finally:
        db.close()


# ── 3. Cobertura: ningún endpoint queda sin clasificar ──────────────────────────────
# Prefijos cubiertos por ATAQUES o LISTAS (se calcula) + rutas que por diseño no llevan id de recurso.
SIN_ID_DE_RECURSO = {
    # escrituras que solo afectan al propio usuario (no reciben ids ajenos)
    "/ahorcado/nueva-partida", "/asistente/mensaje", "/categorias/orden", "/cierre-mes/preferencias",
    "/consentimientos", "/correos/gmail/desconectar", "/correos/gmail/etiqueta", "/correos/gmail/lectura", "/correos/sincronizar",
    "/correos/clave-pdf", "/perfil/punto-de-partida", "/facturas/escanear", "/impuestos/ajustes",
    "/insights/preguntar", "/movimientos/clasificar", "/movimientos/rapido-prueba", "/pagos/checkout",
    "/patrimonio/snapshot", "/patrimonio/activos", "/perfil/api-key/regenerar", "/perfil/recordatorio-diario",
    "/perfil/pais",  # PUT: solo el país del propio usuario (el catálogo y el GET quedan en LISTAS)
    "/perfil/nombre", "/perfil/confianza-autoimport", "/perfil/autoimportar-extractos", "/perfil/eliminar-cuenta",
    "/planificacion/config", "/planificacion/distribuir-ingreso", "/planificacion/plan-meta", "/push/suscribir",
    "/push/desuscribir", "/push/probar", "/reglas-categoria/aprender", "/reglas-categoria", "/categorias",
    "/por-categorizar/silencios/quitar", "/reglas-categoria/sugerencias/crear",
    "/reglas-categoria/sugerencias/descartar", "/reglas-categoria/sugerencias/restaurar",
    "/cuentas", "/presupuestos", "/cierre-mes/{mes}", "/cierre-mes/{mes}/generar", "/gastos-compartidos/contactos", "/cuentas/invitaciones/pendientes",
    "/correos/{correo_id}/aprobar", "/movimientos/rapido-voz", "/resumen/preferencias",
    "/comentarios",  # POST: crea un comentario propio (no recibe ids)
    "/perfil/importar-respaldo", "/perfil/importar-respaldo/vista-previa", "/perfil/importar-respaldo/deshacer",  # crean filas PROPIAS desde un archivo; deshacer solo lote propio
    "/metas/simular",  # solo calcula con los números que envía (más el sobrante propio)
    "/avisos/preferencias", "/avisos/probar", "/avisos/retencion",  # solo administración (404 al resto): no reciben ids ajenos
}
# Rutas sin sesión de usuario por diseño (crons con secreto, callbacks firmados, API key, webhook firmado, públicas).
SIN_SESION = {
    "/backup/cron/enviar", "/cierre-mes/cron/cerrar-mes", "/correos/oauth/callback", "/correos/cron/sincronizar", "/correos/cron/purgar-retencion",
    "/metricas/uso", "/metricas/cron/revisar-uso-semanal", "/metricas/cron/avisos-crecimiento",
    "/avisos/cron/resumen-diario", "/impuestos/cron/avisar-renta", "/costos-bancarios/cron/avisar-cargos",
    "/movimientos/cuentas-rapido", "/movimientos/rapido",
    "/movimientos/rapido-voz", "/pagos/webhook", "/patrimonio/cron/snapshot-mensual",
    "/patrimonio/cron/actualizar-precios", "/push/vapid-public-key", "/push/cron/revisar-presupuestos",
    "/push/cron/aportes-automaticos", "/push/cron/premiar-presupuestos-mes", "/push/cron/resumen-semanal",
    "/push/cron/recordar-registro", "/push/cron/recordar-recurrentes", "/resumen/cron/semanal",
}


def _plantilla_ataque(ruta: str) -> str:
    """Normaliza '/x/{id}/y' para comparar (los nombres de parámetro no importan)."""
    import re

    return re.sub(r"\{[^}]+\}", "{}", ruta.split("?")[0])


def test_todo_endpoint_esta_clasificado():
    cubiertas = {(m, _plantilla_ataque(r)) for m, r, _ in ATAQUES}
    listas = {("GET", _plantilla_ataque(r)) for r in LISTAS}
    sin_id = {_plantilla_ataque(r) for r in SIN_ID_DE_RECURSO}
    sin_sesion = {_plantilla_ataque(r) for r in SIN_SESION}
    sin_clasificar = []
    for metodo, ruta, _usuario, _clave, _mod in RUTAS:
        p = _plantilla_ataque(ruta)
        if (metodo, p) in cubiertas or (metodo, p) in listas or p in sin_id or p in sin_sesion:
            continue
        sin_clasificar.append(f"{metodo} {ruta}")
    assert not sin_clasificar, (
        "Endpoints sin cubrir en test_aislamiento_usuarios.py (agrégalos a ATAQUES si reciben ids de recursos): "
        + ", ".join(sin_clasificar)
    )


# ── 4. Sin sesión ────────────────────────────────────────────────────────────────────
RUTAS_DE_USUARIO = [(m, r) for (m, r, u, k, _) in RUTAS if u]
RUTAS_CRON = [(m, r) for (m, r, u, k, _) in RUTAS if "/cron/" in r or r == "/metricas/uso"]


@pytest.mark.parametrize("metodo,ruta", RUTAS_DE_USUARIO, ids=[f"{m} {r}" for m, r in RUTAS_DE_USUARIO])
def test_sin_sesion_todo_endpoint_de_usuario_da_401(metodo, ruta):
    app.dependency_overrides.pop(get_current_user, None)
    url = ruta.replace("{", "x").replace("}", "")
    r = client.request(metodo, url, json={})
    assert r.status_code == 401, f"{metodo} {ruta} sin sesión -> {r.status_code}"


@pytest.mark.parametrize("metodo,ruta", RUTAS_CRON, ids=[f"{m} {r}" for m, r in RUTAS_CRON])
def test_crons_exigen_secreto(metodo, ruta):
    assert client.request(metodo, ruta).status_code == 401
    assert client.request(metodo, ruta, headers={"X-Cron-Secret": "equivocado"}).status_code == 401
    assert client.request(metodo, ruta, headers={"X-Cron-Secret": ""}).status_code == 401


def test_rutas_con_api_key_rechazan_clave_falsa():
    for metodo, ruta in (("GET", "/movimientos/cuentas-rapido"), ("POST", "/movimientos/rapido"),
                         ("POST", "/movimientos/rapido-voz")):
        r = client.request(metodo, ruta, headers={"X-Api-Key": "cdm_falsa"}, json={"monto": 1, "tipo": "expense",
                                                                                   "descripcion": "x", "texto": "x"})
        assert r.status_code == 401, ruta
        assert client.request(metodo, ruta).status_code in (401, 422)


# ── 5. /consejos: GET sin ids, clasificado como LISTA (B no ve datos de A); aquí además el control ──
def test_consejos_solo_muestran_datos_del_propio_usuario():
    """D tiene una tarjeta al 80 % del cupo (genera un consejo que nombra la tarjeta). B no ve nada de D
    y D sí lo ve: así la prueba de lectura de B no es vacía."""
    from app.derechos_titular import borrar_todo

    D = UsuarioActual(id="aisl-d", email="d-aisl@test.com", nombre="Consejos D")
    assert "/consejos" in LISTAS  # clasificado: lo recorre test_lecturas_de_b_no_muestran_datos_de_a
    try:
        como(D)
        _ok(client.post("/cuentas", json={"nombre": "ZZD Tarjeta", "tipo": "credito", "saldo": -4_000_000,
                                          "cupo_total": 5_000_000, "dia_corte": 10, "dia_pago": 25}))
        propios = _ok(client.get("/consejos"))
        assert "ZZD Tarjeta" in json.dumps(propios) and propios["aviso"]
        como(B)
        ajeno = client.get("/consejos")
        assert ajeno.status_code == 200 and "zzd" not in ajeno.text.lower()
        como(A)
        assert "zzd" not in client.get("/consejos").text.lower()
    finally:
        db = SessionLocal()
        try:
            borrar_todo(db, D.id, D.email)
        finally:
            db.close()


def test_preferencias_de_lectura_del_correo_son_por_usuario():
    """Apagar un uso del correo (movimientos/extractos) solo afecta a quien lo apaga."""
    _sembrar()
    como(B)
    assert client.put("/correos/gmail/lectura", json={"leer_extractos": False}).status_code == 200
    assert client.get("/correos/gmail/estado").json()["leer_extractos"] is False
    como(A)
    assert client.get("/correos/gmail/estado").json()["leer_extractos"] is True
    carriles = client.get("/correos/carriles").json()
    assert carriles["extractos"]["activo"] is True
    como(B)
    assert client.get("/correos/carriles").json()["extractos"]["activo"] is False

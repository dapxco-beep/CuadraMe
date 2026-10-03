"""Impuestos Colombia (4x1000 y posibles deducibles) y Cierre de mes automático.
Los números esperados están calculados a mano (ver comentarios)."""
import io
import os
import tempfile
from datetime import datetime

os.environ.setdefault("DATABASE_URL", f"sqlite:///{os.path.join(tempfile.mkdtemp(), 'imp_cierre.db')}")

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.auth import UsuarioActual, get_current_user  # noqa: E402
from app.cierre_mensual import categoria_que_mas_subio, generar_cierre  # noqa: E402
from app.impuestos import AVISO, calcular_gmf, organizar_deducibles, sugerir_categoria_fiscal, uvt_del_anio  # noqa: E402
from app.main import app  # noqa: E402
from app.routers import cierre as cierre_mod  # noqa: E402
from app.routers.cierre import mes_anterior  # noqa: E402
from app.utiles import ahora_co  # noqa: E402

client = TestClient(app)
CABECERA = {"X-Cron-Secret": "secreto-de-prueba"}
USUARIO = {"actual": UsuarioActual(id="ic-0", email="ic@test.com", nombre="IC")}


@pytest.fixture(autouse=True)
def _ambiente():
    app.dependency_overrides[get_current_user] = lambda: USUARIO["actual"]
    yield


def como(uid, email=None):
    USUARIO["actual"] = UsuarioActual(id=uid, email=email or f"{uid}@test.com", nombre=uid)


# ── 4x1000 (puro) ────────────────────────────────────

def _g(monto, desc="Compra", cuenta="A", tipo="expense", fuente="manual", notas="", fecha=datetime(2026, 10, 5)):
    return {"fecha": fecha, "monto": monto, "tipo": tipo, "fuente": fuente, "notas": notas,
            "descripcion": desc, "cuenta_id": cuenta}


def test_uvt_conocida_y_fallback():
    assert uvt_del_anio(2026) == 52374
    assert uvt_del_anio(2030) == 52374   # sin dato: usa la última conocida


def test_gmf_cuenta_normal_y_exenta():
    # UVT 2026 = 52.374 -> 350 UVT = 18.330.900 de débitos mensuales exentos
    tipos = {"A": "ahorro", "E": "ahorro", "T": "credito", "C": "efectivo"}
    movs = [
        _g(10_000_000, "Arriendo y mercado", "A"),
        _g(5_000_000, "Pago tarjeta Visa", "A"),                       # sale de la cuenta: SÍ es débito gravable
        _g(60_000, "GRAVAMEN MOVIMIENTOS FINANCIEROS", "A"),           # el cobro del GMF no se grava a sí mismo
        _g(1_000_000, "Transferencia a Nequi", "A", "transfer", "transferencia_interna", "TRF:g1:out"),
        _g(3_000_000, "Nómina", "A", "income"),                        # entra: no es débito
        _g(20_000_000, "Compra de carro", "E"),
        _g(900_000, "Compra con tarjeta", "T"),                        # tarjeta de crédito: no es cuenta bancaria
        _g(400_000, "Mercado en efectivo", "C"),                       # efectivo: sin GMF
    ]
    r = calcular_gmf(movs, tipos, "E", 2026)
    assert len(r["meses"]) == 1
    m = r["meses"][0]
    # A: 10.000.000 + 5.000.000 + 1.000.000 = 16.000.000 | E: 20.000.000 - 18.330.900 = 1.669.100
    assert m["debitos"] == 36_000_000
    assert m["base_gravable"] == 17_669_100
    assert m["exento_usado"] == 18_330_900
    assert m["gmf_estimado"] == 70_676          # 17.669.100 * 0,004 = 70.676,4
    assert m["gmf_registrado"] == 60_000
    assert r["total_estimado"] == 70_676 and r["tope_exento_mensual"] == 18_330_900


def test_gmf_cuenta_exenta_bajo_el_tope_no_paga_y_sin_marca_si_paga():
    tipos = {"E": "ahorro"}
    movs = [_g(10_000_000, cuenta="E")]
    assert calcular_gmf(movs, tipos, "E", 2026)["total_estimado"] == 0
    assert calcular_gmf(movs, tipos, None, 2026)["total_estimado"] == 40_000   # 10.000.000 * 0,004


def test_gmf_separa_meses_y_anios():
    tipos = {"A": "ahorro"}
    movs = [_g(1_000_000, fecha=datetime(2026, 9, 1)), _g(2_000_000, fecha=datetime(2026, 10, 1)),
            _g(9_000_000, fecha=datetime(2025, 10, 1))]
    r = calcular_gmf(movs, tipos, None, 2026)
    assert [(f["mes"], f["gmf_estimado"]) for f in r["meses"]] == [("2026-09", 4_000), ("2026-10", 8_000)]
    assert r["promedio_mensual"] == 6_000


# ── Deducibles (puro) ───────────────────────────────

def _gasto(id, desc, monto, cat="Otros", fecha=datetime(2026, 5, 1)):
    return {"id": id, "descripcion": desc, "monto": monto, "categoria": cat, "fecha": fecha}


def test_sugerencias_por_descripcion_y_categoria():
    assert sugerir_categoria_fiscal(_gasto("1", "Colsanitas medicina prepagada", 1)) == ("salud", "alta")
    assert sugerir_categoria_fiscal(_gasto("2", "Intereses crédito de vivienda", 1)) == ("vivienda_intereses", "alta")
    assert sugerir_categoria_fiscal(_gasto("3", "Cuota ICETEX", 1)) == ("educacion", "alta")
    assert sugerir_categoria_fiscal(_gasto("4", "Consulta", 1, cat="Salud")) == ("salud", "baja")
    assert sugerir_categoria_fiscal(_gasto("5", "GRAVAMEN MOVIMIENTOS FINANCIEROS", 1)) == ("gmf", "alta")
    assert sugerir_categoria_fiscal(_gasto("6", "Cine con amigos", 1, cat="Diversión")) is None


def test_organizar_deducibles_solo_suma_como_confirmado_lo_que_el_usuario_reviso():
    gastos = [
        _gasto("a", "Colsanitas medicina prepagada", 300_000),
        _gasto("b", "Intereses crédito de vivienda", 450_000),
        _gasto("c", "Matrícula Universidad X", 2_000_000),
        _gasto("d", "Consulta médica", 80_000, cat="Salud"),
        _gasto("e", "Cine", 40_000, cat="Diversión"),
        _gasto("f", "Donación Cruz Roja", 50_000),
        _gasto("g", "GRAVAMEN MOVIMIENTOS FINANCIEROS", 60_000),
    ]
    marcas = {"c": {"categoria_fiscal": "educacion", "soporte": "Recibo en el correo, feb 2026", "estado": "confirmada"},
              "d": {"categoria_fiscal": "salud", "soporte": "", "estado": "descartada"}}
    r = organizar_deducibles(gastos, marcas)
    ids = [i["movimiento_id"] for i in r["items"]]
    assert "d" not in ids and "e" not in ids                       # descartado y sin relación fiscal
    assert r["total_confirmado"] == 2_000_000                       # solo la matrícula, que el usuario confirmó
    assert r["total_sugerido"] == 300_000 + 450_000 + 50_000 + 60_000
    edu = [c for c in r["por_categoria"] if c["categoria_fiscal"] == "educacion"][0]
    assert edu["confirmado"] == 2_000_000 and edu["con_soporte"] == 1
    c = [i for i in r["items"] if i["movimiento_id"] == "c"][0]
    assert c["estado"] == "confirmado" and c["soporte"].startswith("Recibo")


# ── Impuestos (API) ─────────────────────────────────

def _usuario_con_gastos(uid):
    como(uid)
    ahorro = client.post("/cuentas", json={"nombre": "Ahorros", "saldo": 0}).json()["id"]
    visa = client.post("/cuentas", json={"nombre": "Visa", "tipo": "credito"}).json()["id"]

    def mov(cuenta, tipo, monto, desc, cat=None):
        b = {"cuenta_id": cuenta, "tipo": tipo, "monto": monto, "descripcion": desc}
        if cat:
            b["categoria"] = cat
        return client.post("/movimientos", json=b).json()["id"]

    mov(ahorro, "income", 5_000_000, "Nómina", "Ingresos")
    prepagada = mov(ahorro, "expense", 300_000, "Colsanitas medicina prepagada", "Salud")
    mov(ahorro, "expense", 1_000_000, "Mercado", "Comida")
    mov(visa, "expense", 200_000, "Restaurante", "Comida")
    mov(ahorro, "expense", 200_000, "Pago tarjeta Visa")
    return ahorro, visa, prepagada


def test_resumen_impuestos_gmf_y_deducibles():
    ahorro, visa, prepagada = _usuario_con_gastos("imp-api")
    j = client.get("/impuestos/resumen").json()
    assert j["aviso"] == AVISO and "no asesoría" in j["aviso"]
    # débitos de la cuenta de ahorro: 300.000 + 1.000.000 + 200.000 (pago de tarjeta) = 1.500.000
    # (la compra en la tarjeta de crédito no es débito bancario) -> 1.500.000 * 0,004 = 6.000
    assert j["gmf"]["total_estimado"] == 6_000
    assert [c["id"] for c in j["cuentas_ahorro"]] == [ahorro]
    items = j["deducibles"]["items"]
    assert [(i["movimiento_id"], i["categoria_fiscal"], i["estado"]) for i in items] == [(prepagada, "salud", "sugerido")]
    assert j["deducibles"]["total_confirmado"] == 0     # nada se promete sin que el usuario lo confirme


def test_cuenta_exenta_baja_el_4x1000_y_solo_acepta_ahorros_propios():
    ahorro, visa, _ = _usuario_con_gastos("imp-exenta")
    assert client.put("/impuestos/ajustes", json={"cuenta_exenta_gmf_id": visa}).status_code == 404   # tarjeta: no
    assert client.put("/impuestos/ajustes", json={"cuenta_exenta_gmf_id": ahorro}).status_code == 200
    j = client.get("/impuestos/resumen").json()
    assert j["cuenta_exenta_gmf_id"] == ahorro and j["gmf"]["total_estimado"] == 0   # 1.500.000 < 18.330.900
    client.put("/impuestos/ajustes", json={"cuenta_exenta_gmf_id": None})
    assert client.get("/impuestos/resumen").json()["gmf"]["total_estimado"] == 6_000
    como("imp-otro")
    assert client.put("/impuestos/ajustes", json={"cuenta_exenta_gmf_id": ahorro}).status_code == 404


def test_marcar_confirmar_soporte_descartar_y_quitar():
    _, _, prepagada = _usuario_con_gastos("imp-marcas")
    r = client.put(f"/impuestos/marcas/{prepagada}", json={"categoria_fiscal": "salud", "soporte": "Certificado 2026 en Drive"})
    assert r.status_code == 200
    d = client.get("/impuestos/resumen").json()["deducibles"]
    assert d["total_confirmado"] == 300_000 and d["items"][0]["soporte"] == "Certificado 2026 en Drive"
    client.put(f"/impuestos/marcas/{prepagada}", json={"categoria_fiscal": "salud", "estado": "descartada"})
    assert client.get("/impuestos/resumen").json()["deducibles"]["items"] == []
    client.delete(f"/impuestos/marcas/{prepagada}")   # vuelve a ser sugerencia
    assert client.get("/impuestos/resumen").json()["deducibles"]["items"][0]["estado"] == "sugerido"
    assert client.put(f"/impuestos/marcas/{prepagada}", json={"categoria_fiscal": "inventada"}).status_code == 422
    como("imp-intruso")
    assert client.put(f"/impuestos/marcas/{prepagada}", json={"categoria_fiscal": "salud"}).status_code == 404


def test_exportar_excel_trae_aviso_resumen_y_detalle():
    _, _, prepagada = _usuario_con_gastos("imp-xlsx")
    client.put(f"/impuestos/marcas/{prepagada}", json={"categoria_fiscal": "salud", "soporte": "Factura 123"})
    r = client.get("/impuestos/exportar.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Resumen", "4x1000 por mes", "Posibles deducibles"]
    assert "no asesoría" in wb["Resumen"]["A2"].value
    filas = list(wb["Posibles deducibles"].iter_rows(values_only=True))
    assert filas[1][1] == "Colsanitas medicina prepagada" and filas[1][4] == "Confirmado por ti" and filas[1][5] == "Factura 123"
    assert [f for f in wb["4x1000 por mes"].iter_rows(values_only=True)][1][4] == 6_000


def test_anio_fuera_de_rango():
    como("imp-anio")
    assert client.get("/impuestos/resumen?anio=1999").status_code == 422


# ── Cierre de mes (puro) ──────────────────────────────

def _f(tipo, monto, cat="Otros", desc="x"):
    return {"tipo": tipo, "monto": monto, "categoria": cat, "descripcion": desc}


def test_cierre_numeros_a_mano():
    mes = [_f("income", 3_000_000, "Ingresos"), _f("expense", 400_000, "Comida"), _f("expense", 200_000, "Transporte"),
           _f("expense", 80_000, "Intereses", "Intereses de financiación"), _f("expense", 100_000, "Ocio")]
    previo = [_f("expense", 300_000, "Comida"), _f("expense", 200_000, "Transporte"), _f("expense", 20_000, "Ocio")]
    r = generar_cierre("2026-09", mes, previo, intereses_estimados_deudas=45_000.4,
                       patrimonio_inicio=10_000_000, patrimonio_fin=10_500_000)
    assert (r["ingresos"], r["gastos"], r["ahorro"], r["tasa_ahorro"]) == (3_000_000, 780_000, 2_220_000, 74.0)
    assert r["gasto_mes_anterior"] == 520_000
    # Comida +100.000 (33,3 %), Ocio +80.000 (400 %), Intereses +80.000 (nueva): gana el mayor aumento ABSOLUTO
    assert r["categoria_que_mas_subio"] == {"categoria": "Comida", "monto": 400_000, "monto_anterior": 300_000,
                                            "aumento": 100_000, "variacion_pct": 33.3}
    assert r["intereses_pagados"] == 80_000 and r["intereses_estimados_deudas"] == 45_000
    assert r["patrimonio"] == {"inicio": 10_000_000, "fin": 10_500_000, "variacion": 500_000}
    assert [c["categoria"] for c in r["top_categorias"]] == ["Comida", "Transporte", "Ocio"]
    assert len(r["acciones"]) == 2 and "Comida" in r["acciones"][0] and "intereses" in r["acciones"][1]
    assert r["titulo"] == "Cierre de septiembre de 2026" and not r["sin_datos"]


def test_cierre_gastos_mayores_que_ingresos_y_sin_datos():
    r = generar_cierre("2026-01", [_f("income", 1_000_000), _f("expense", 1_300_000, "Comida")], [])
    assert r["ahorro"] == -300_000 and "te pasaste $300.000" in r["resumen_corto"]
    assert "$300.000" in r["acciones"][0] and "Comida" in r["acciones"][0]
    vacio = generar_cierre("2026-01", [], [])
    assert vacio["sin_datos"] and vacio["patrimonio"] is None and vacio["acciones"]


def test_alza_pequena_no_cuenta():
    assert categoria_que_mas_subio({"A": 110_000}, {"A": 100_000}) is None        # +10 %
    assert categoria_que_mas_subio({"A": 120_000}, {"A": 100_000}) is None        # +20 % pero solo +20.000 (< 30.000)
    assert categoria_que_mas_subio({"A": 240_000}, {"A": 200_000})["variacion_pct"] == 20.0
    assert categoria_que_mas_subio({"A": 10_000}, {}) is None                      # nueva pero < 30.000
    assert categoria_que_mas_subio({"A": 50_000}, {})["variacion_pct"] is None     # nueva, sin % definido


def test_mes_anterior():
    assert mes_anterior("2026-01") == "2025-12" and mes_anterior("2026-10") == "2026-09"


# ── Cierre de mes (API + cron) ────────────────────────

def _mes_pasado() -> str:
    return mes_anterior(ahora_co().strftime("%Y-%m"))


def _datos_mes_pasado(uid, email=None):
    como(uid, email)
    cid = client.post("/cuentas", json={"nombre": "Banco"}).json()["id"]
    mes = _mes_pasado()

    def mov(tipo, monto, desc, cat, dia):
        client.post("/movimientos", json={"cuenta_id": cid, "tipo": tipo, "monto": monto, "descripcion": desc,
                                           "categoria": cat, "fecha": f"{mes}-{dia:02d}T12:00:00"})

    mov("income", 2_000_000, "Nómina", "Ingresos", 5)
    mov("expense", 500_000, "Mercado", "Comida", 10)
    mov("expense", 100_000, "Pago tarjeta Visa", "Deudas", 15)    # no cuenta como gasto
    return mes


def test_cron_cierre_exige_secreto_y_mes_terminado():
    assert client.post("/cierre-mes/cron/cerrar-mes").status_code == 401
    assert client.post("/cierre-mes/cron/cerrar-mes", headers={"X-Cron-Secret": "mal"}).status_code == 401
    actual = ahora_co().strftime("%Y-%m")
    assert client.post(f"/cierre-mes/cron/cerrar-mes?mes={actual}", headers=CABECERA).status_code == 400
    assert client.post("/cierre-mes/cron/cerrar-mes?mes=2026-13", headers=CABECERA).status_code == 422


def test_cron_genera_guarda_es_idempotente_y_usa_la_regla_contable(monkeypatch):
    correos = []
    monkeypatch.setattr(cierre_mod, "_enviar_correo", lambda destino, asunto, texto: correos.append(destino) or True)
    mes = _datos_mes_pasado("cierre-a", "cierre-a@test.com")
    client.put("/cierre-mes/preferencias", json={"correo": True})
    _datos_mes_pasado("cierre-b", "cierre-b@test.com")   # NO activó el correo

    r1 = client.post("/cierre-mes/cron/cerrar-mes", headers=CABECERA).json()
    assert r1["ok"] and r1["mes"] == mes and r1["generados"] >= 2 and r1["errores"] == 0

    como("cierre-a", "cierre-a@test.com")
    c = client.get(f"/cierre-mes/{mes}").json()
    assert (c["ingresos"], c["gastos"], c["ahorro"]) == (2_000_000, 500_000, 1_500_000)   # el pago de tarjeta no cuenta
    assert correos == ["cierre-a@test.com"]                                              # solo quien lo activó

    r2 = client.post("/cierre-mes/cron/cerrar-mes", headers=CABECERA).json()
    assert r2["generados"] == 0 and r2["ya_existian"] >= 2 and r2["push"] == 0 and r2["correos"] == 0
    assert correos == ["cierre-a@test.com"]                                              # no se repite el correo
    assert c == client.get(f"/cierre-mes/{mes}").json()

    lista = client.get("/cierre-mes").json()
    assert [x["mes"] for x in lista] == [mes] and lista[0]["visto"] is True


def test_preferencia_de_correo_es_opt_in_y_se_puede_apagar():
    como("cierre-pref", "pref@test.com")
    assert client.get("/cierre-mes/preferencias").json()["correo"] is False
    client.put("/cierre-mes/preferencias", json={"correo": True})
    assert client.get("/cierre-mes/preferencias").json()["correo"] is True
    client.put("/cierre-mes/preferencias", json={"correo": False})
    assert client.get("/cierre-mes/preferencias").json()["correo"] is False


def test_generar_manual_mes_pasado_parcial_y_errores():
    mes = _datos_mes_pasado("cierre-man")
    assert client.get(f"/cierre-mes/{mes}").status_code == 404
    g1 = client.post(f"/cierre-mes/{mes}/generar").json()
    g2 = client.post(f"/cierre-mes/{mes}/generar").json()
    assert g1 == g2 and g1["ingresos"] == 2_000_000
    actual = ahora_co().strftime("%Y-%m")
    parcial = client.post(f"/cierre-mes/{actual}/generar").json()
    assert parcial["parcial"] is True
    assert client.get(f"/cierre-mes/{actual}").status_code == 404      # el parcial no se guarda
    assert client.post("/cierre-mes/2099-01/generar").status_code == 400
    assert client.post("/cierre-mes/2020-01/generar").status_code == 404   # sin movimientos
    como("cierre-ajeno")
    assert client.get(f"/cierre-mes/{mes}").status_code == 404           # nadie ve el cierre de otro

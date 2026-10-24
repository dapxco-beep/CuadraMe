"""El Excel de movimientos viene listo para usar: formatos reales, español, cuenta, filtros, resumen y cuentas.
Datos 100 % inventados."""
import io
import os
import tempfile
from datetime import datetime

os.environ.setdefault("DATABASE_URL", f"sqlite:///{os.path.join(tempfile.mkdtemp(), 'exportar_excel.db')}")

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.auth import UsuarioActual, get_current_user  # noqa: E402
from app.database import Base, SessionLocal, engine  # noqa: E402
from app.main import app  # noqa: E402
from app.models import Cuenta, Movimiento  # noqa: E402

Base.metadata.create_all(bind=engine)
client = TestClient(app)


def _libro():
    uid = "xl-1"
    app.dependency_overrides[get_current_user] = lambda: UsuarioActual(id=uid, email="x@t.com", nombre="Persona")
    if getattr(_libro, "listo", False):
        return load_workbook(io.BytesIO(client.get("/exportar/movimientos.xlsx").content))
    _libro.listo = True
    db = SessionLocal()
    try:
        ahorro = Cuenta(usuario_id=uid, nombre="Ahorros Ejemplo", tipo="ahorro", saldo=900_000.0)
        tarjeta = Cuenta(usuario_id=uid, nombre="Tarjeta Ejemplo", tipo="credito", saldo=-100_000.0, cupo_total=1_000_000.0)
        db.add_all([ahorro, tarjeta])
        db.flush()
        for cuenta, dia, desc, monto, tipo, cat in [
            (ahorro, 3, "Nomina ficticia", 1_000_000.0, "income", "Ingresos"),
            (ahorro, 9, "Mercado ficticio", 60_000.5, "expense", "Comida"),
            (tarjeta, 10, "Compra ficticia", 100_000.0, "expense", "Hogar"),
        ]:
            db.add(Movimiento(usuario_id=uid, cuenta_id=cuenta.id, descripcion=desc, categoria=cat, tipo=tipo, monto=monto,
                              fecha=datetime(2026, 9, dia), fuente="manual"))
        db.commit()
    finally:
        db.close()
    r = client.get("/exportar/movimientos.xlsx")
    assert r.status_code == 200
    return load_workbook(io.BytesIO(r.content))


def test_el_excel_trae_formatos_reales_espanol_y_cuenta():
    wb = _libro()
    assert wb.sheetnames == ["Movimientos", "Resumen", "Cuentas"]
    ws = wb["Movimientos"]
    enc = [c.value for c in ws[1]]
    fila = {k: c for k, c in zip(enc, ws[2])}
    assert isinstance(fila["Fecha"].value, datetime) and fila["Fecha"].number_format == "dd/mm/yyyy"  # fecha real, no texto
    assert fila["Tipo"].value == "Ingreso" and fila["Cuenta"].value == "Ahorros Ejemplo"
    assert "$" in fila["Monto"].number_format and isinstance(fila["Monto"].value, (int, float))
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref  # primera fila fija y filtros
    gasto = {k: c.value for k, c in zip(enc, ws[3])}
    assert gasto["Valor (+ entra / − sale)"] == -60_000.5  # lo que sale va en negativo


def test_el_resumen_suma_con_la_regla_de_la_app_y_las_cuentas_muestran_deuda():
    wb = _libro()
    rs = wb["Resumen"]
    filas = [[c.value for c in f] for f in rs.iter_rows(min_row=2)]
    sep = next(f for f in filas if f[0] == "Septiembre 2026")
    assert sep[1] == 1_000_000.0 and sep[2] == 160_000.5 and sep[3] == 839_999.5
    assert ["Hogar", 100_000.0] == next(f[:2] for f in filas if f[0] == "Hogar")
    cs = wb["Cuentas"]
    cuentas = {f[0].value: [c.value for c in f] for f in cs.iter_rows(min_row=2)}
    assert cuentas["Tarjeta Ejemplo"][2] == 100_000.0 and cuentas["Tarjeta Ejemplo"][3] == 1_000_000.0  # deuda positiva y su cupo
    assert cuentas["Ahorros Ejemplo"][2] == 900_000.0
